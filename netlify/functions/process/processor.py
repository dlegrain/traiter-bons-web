"""
processor.py - In-memory version of traiter_bons for the Netlify function.
Interactive parts removed: unknown codes get 'UNK' and are listed in a text file.
"""
import re
import io
import os
import zipfile
import fitz
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SITES_ALDI = {
    "1": "ALDI N.V. - ZEMST",
    "2": "ALDI N.V. - ERPE MERE",
    "3": "ALDI N.V. - HEUSDEN-ZOLDER",
}

KEYWORD_MAP = [
    ("lift: dieptereiniging",                           "72"),
    ("ascenseur : nettoyage en profondeur",             "72"),
    ("winkelkarren en winkelmandje",                    "79"),
    ("caddys et paniers",                               "79"),
    ("traphal",                                         "07"),
    ("cage d'escalier",                                 "07"),
    ("roltrap / looppad: dieptereiniging",              "21"),
    ("escalator",                                       "21"),
    ("roltrap",                                         "21"),
    ("looppad",                                         "21"),
    ("reinigen complete inkomsas",                      "14"),
    ("nettoyage complet du sas",                        "14"),
    ("parkeerplaatsen & caddyzone",                     "01"),
    ("schoonmaak van parkeerplaatsen",                  "01"),
    ("nettoyage de parking & zone caddy",               "01"),
    ("parkeerplaatsen",                                 "01"),
    ("parking & zone caddy",                            "01"),
    ("reinigen spiegels op kolommen",                   "74"),
    ("nettoyage des miroirs",                           "74"),
    ("laadkade",                                        "09"),
    ("quai de d",                                       "09"),
    ("quai de chargement",                              "09"),
    ("onderdekte parking",                              "78"),
    ("parking couvert",                                 "78"),
    ("filiaalspecifieke taken",                         "90"),
    ("filiaalspecifieke",                               "90"),
    ("taches specifiques",                              "90"),
    ("evacuation d'encombrants",                        "UNK"),
    ("afvoeren van afval",                              "UNK"),
    ("nettoyage regulier des locaux",                   "UNK"),
    ("regelmatige schoonmaak van lokalen",              "UNK"),
    ("inkommat: stofzuigen",                            "15"),
    ("inkommat",                                        "15"),
    ("spinnenwebben en vuil weghalen",                  "23"),
    ("spinnenwebben",                                   "23"),
    ("toiles d'araign",                                 "23"),
    ("vitrinekast",                                     "70"),
    ("vitrine",                                         "70"),
]

MAPPING_PRESTATIONS = [
    ("01", "Nettoyage de Parking & zone caddy / Schoonmaak van parkeerplaatsen & caddyzone"),
    ("07", "Cage d'escalier / Traphal"),
    ("09", "Nettoyage du quai de déchargement / Laadkade"),
    ("14", "Nettoyage complet du sas d'entrée / Reinigen complete inkomsas"),
    ("15", "Tapis à l'entrée : Aspirer / Inkommat: Stofzuigen"),
    ("21", "Escalator / tapis roulant : nettoyage en profondeur / Roltrap / looppad: Dieptereiniging"),
    ("23", "Toiles d'araignée : locaux techniques / Spinnenwebben en vuil weghalen"),
    ("70", "Vitrine : intérieur et extérieur / Vitrinekast binnen- en buitenkant"),
    ("72", "Ascenseur : nettoyage en profondeur / Lift: dieptereiniging"),
    ("74", "Nettoyage des miroirs sur les colonnes / Reinigen spiegels op kolommen"),
    ("78", "Parking couvert : nettoyage humide / Onderdekte parking: natreiniging"),
    ("79", "Caddys et paniers : nettoyage à l'eau / Winkelkarren en winkelmandje: natreiniging"),
    ("90", "Tâches spécifiques aux succursales / Filiaalspecifieke taken"),
]


def extraire_filiale(texte: str) -> str:
    m = re.search(r'SDIC\d+\s*-\s*BE\d+\s*-\s*(.+)', texte)
    return m.group(1).strip() if m else ""


def extraire_ligne_prestation(texte: str) -> str:
    lignes = [l.strip() for l in texte.split("\n") if l.strip()]
    cdic_idx = next(
        (i for i, l in enumerate(lignes) if re.match(r"CDIC\d+", l)),
        len(lignes)
    )
    return " ".join(lignes[min(4, cdic_idx):cdic_idx])


def extraire_code(texte: str, unknown_codes: list) -> str:
    m = re.search(r'\b(\d{2})\.\s?[A-Za-z]', texte)
    if m:
        return m.group(1)

    libelle = extraire_ligne_prestation(texte)
    libelle_lower = libelle.lower()
    for keyword, code in KEYWORD_MAP:
        if keyword.lower() in libelle_lower:
            return code

    unknown_codes.append(libelle or texte[:120])
    return "UNK"


def extraire_metadonnees_page(texte: str, unknown_codes: list):
    m_site = re.search(r'\b(BE\d{6})\b', texte)
    site = m_site.group(1) if m_site else "UNKNOWN"

    m_date = re.search(
        r'Mission\s+d[ée]but[ée]e?\s+le\s+(\d{2}/\d{2}/\d{4})',
        texte, re.IGNORECASE
    )
    if m_date:
        d, mo, y = m_date.group(1).split("/")
        date_str = f"{d}-{mo}-{y}"
    else:
        date_str = "00-00-0000"

    code = extraire_code(texte, unknown_codes)
    return site, date_str, code


def has_client_signature(doc: fitz.Document) -> bool:
    for page in doc:
        mid_x = page.rect.width / 2
        for drawing in page.get_drawings():
            if len(drawing.get("items", [])) <= 5:
                continue
            rect = drawing.get("rect")
            if rect:
                centre_x = (rect.x0 + rect.x1) / 2
                if centre_x >= mid_x:
                    return True
    return False


def supprimer_note(doc: fitz.Document):
    for page in doc:
        rects = page.search_for("Note attribu")
        for rect in rects:
            zone = fitz.Rect(rect.x0, rect.y0, page.rect.width - 36, rect.y1)
            page.add_redact_annot(zone, fill=(1, 1, 1))
        if rects:
            page.apply_redactions()


def extraire_technicien_from_text(texte: str) -> str:
    lignes = [l.strip() for l in texte.splitlines() if l.strip()]
    prefixes_a_ignorer = (
        "CDIC", "SDIC", "Adresse", "Mission", "Horaire",
        "Commentaire", "Signature", "Note", "Rapport",
        "Powered", "Edit", "Edité"
    )
    for ligne in lignes:
        if re.match(r"^\d[\d\s]*-\s*.+", ligne):
            technicien = re.sub(r"^\d[\d\s]*-\s*", "", ligne).strip()
            if technicien:
                return technicien
    for ligne in lignes:
        if " - " in ligne and not ligne.startswith(prefixes_a_ignorer):
            return ligne.strip()
    for i, ligne in enumerate(lignes):
        if ligne.startswith("CDIC") and i > 0:
            prev = lignes[i - 1]
            if not prev.startswith(prefixes_a_ignorer):
                return prev.strip()
    return "Non trouvé"


def charger_pointages_filtres(excel_bytes: bytes, site_nom: str):
    df = pd.read_excel(io.BytesIO(excel_bytes), dtype=str)
    col_client = df.columns[0]
    col_debut_pointe = df.columns[6]

    df_aldi = df[df[col_client].str.contains("ALDI", case=False, na=False)]
    if df_aldi.empty:
        return None

    df_site = df_aldi[df_aldi[col_client].str.strip() == site_nom]
    if df_site.empty:
        return None

    masque_vide = df_site[col_debut_pointe].isna() | (df_site[col_debut_pointe].str.strip() == "")
    df_filtre = df_site[masque_vide]
    return df_filtre if not df_filtre.empty else None


def _make_header_cell(ws, row, col, value, header_font, header_fill, border):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = border
    return cell


def generer_excel(signes_pdfs: dict, non_signes_pdfs: dict, df_pointages=None) -> bytes:
    wb = Workbook()
    ws_signes = wb.active
    ws_signes.title = "Signes"
    ws_non_signes = wb.create_sheet("Non_signes")
    ws_recap = wb.create_sheet("Recapitulatif")
    ws_ns_sd = wb.create_sheet("Non_signes_sans_doublons")

    signes = sorted(signes_pdfs.keys())
    non_signes = sorted(non_signes_pdfs.keys())

    fill_blue = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    fill_dark = PatternFill("solid", fgColor="FF2E4057")
    fill_doublon = PatternFill("solid", fgColor="FFFFC7CE")
    hdr_font = Font(bold=True, color="FFFFFFFF", name="Arial", size=10)
    hdr_font_simple = Font(color="FFFFFF", bold=True)
    data_font = Font(name="Arial", size=10)
    thin = Side(style="thin", color="FFBDBDBD")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    for ws, rows, label in [
        (ws_signes, signes, "Signe"),
        (ws_non_signes, non_signes, "Non signe"),
    ]:
        ws.append(["Nom du fichier", "Statut"])
        for nom in rows:
            ws.append([nom, label])
        ws["A1"].fill = fill_blue
        ws["A1"].font = hdr_font_simple
        ws["B1"].fill = fill_blue
        ws["B1"].font = hdr_font_simple
        ws.column_dimensions["A"].width = 55
        ws.column_dimensions["B"].width = 18
        ws.freeze_panes = "A2"

    for col_idx, header in enumerate(["Nom du fichier", "Statut", "Nom raccourci", "Technicien"], start=1):
        _make_header_cell(ws_recap, 1, col_idx, header, hdr_font, fill_dark, border)
    ws_recap.row_dimensions[1].height = 18

    recap_rows = [(n, "Signe", signes_pdfs[n]) for n in signes]
    recap_rows += [(n, "Non signe", non_signes_pdfs[n]) for n in non_signes]

    raccourci_counts = {}
    for nom_f, _, _ in recap_rows:
        k = nom_f[:22]
        raccourci_counts[k] = raccourci_counts.get(k, 0) + 1

    for r_idx, (nom_f, statut, pdf_b) in enumerate(recap_rows, start=2):
        nom_raccourci = nom_f[:22]
        try:
            doc = fitz.open(stream=pdf_b, filetype="pdf")
            technicien = extraire_technicien_from_text(doc[0].get_text()) if len(doc) > 0 else "Non trouvé"
            doc.close()
        except Exception:
            technicien = "Non trouvé"

        for c_idx, value in enumerate([nom_f, statut, nom_raccourci, technicien], start=1):
            cell = ws_recap.cell(row=r_idx, column=c_idx, value=value)
            cell.font = data_font
            cell.border = border
            cell.alignment = Alignment(vertical="center")
        if raccourci_counts.get(nom_raccourci, 0) == 1:
            ws_recap.cell(row=r_idx, column=3).fill = fill_doublon

    for col_idx, width in enumerate([35, 12, 25, 25], start=1):
        ws_recap.column_dimensions[ws_recap.cell(row=1, column=col_idx).column_letter].width = width
    ws_recap.freeze_panes = "A2"

    for col_idx, header in enumerate(["Nom du fichier", "Statut", "Nom raccourci", "Technicien", "Filiale"], start=1):
        _make_header_cell(ws_ns_sd, 1, col_idx, header, hdr_font, fill_dark, border)
    ws_ns_sd.row_dimensions[1].height = 18

    ns_row = 2
    for nom_f in non_signes:
        if raccourci_counts.get(nom_f[:22], 0) != 1:
            continue
        pdf_b = non_signes_pdfs[nom_f]
        try:
            doc = fitz.open(stream=pdf_b, filetype="pdf")
            if len(doc) > 0:
                texte = doc[0].get_text()
                technicien = extraire_technicien_from_text(texte)
                filiale = extraire_filiale(texte)
            else:
                technicien, filiale = "Non trouvé", ""
            doc.close()
        except Exception:
            technicien, filiale = "Non trouvé", ""

        for c_idx, value in enumerate([nom_f, "Non signe", nom_f[:22], technicien, filiale], start=1):
            cell = ws_ns_sd.cell(row=ns_row, column=c_idx, value=value)
            cell.font = data_font
            cell.border = border
            cell.alignment = Alignment(vertical="center")
        ws_ns_sd.cell(row=ns_row, column=3).fill = fill_doublon
        ns_row += 1

    for col_idx, width in enumerate([35, 12, 25, 25, 20], start=1):
        ws_ns_sd.column_dimensions[ws_ns_sd.cell(row=1, column=col_idx).column_letter].width = width
    ws_ns_sd.freeze_panes = "A2"

    if df_pointages is not None and not df_pointages.empty:
        ws_pt = wb.create_sheet("Pointages_non_pointes")
        for col_idx, entete in enumerate(df_pointages.columns, start=1):
            _make_header_cell(ws_pt, 1, col_idx, entete, hdr_font, fill_dark, border)
        ws_pt.row_dimensions[1].height = 18
        for r_idx, row_data in enumerate(df_pointages.itertuples(index=False), start=2):
            for c_idx, value in enumerate(row_data, start=1):
                val = None if (isinstance(value, float) and pd.isna(value)) else value
                cell = ws_pt.cell(row=r_idx, column=c_idx, value=val)
                cell.font = data_font
                cell.border = border
                cell.alignment = Alignment(vertical="center")
        for col in ws_pt.columns:
            max_len = max((len(str(c.value)) if c.value else 0) for c in col)
            ws_pt.column_dimensions[col[0].column_letter].width = min(max_len + 4, 50)
        ws_pt.freeze_panes = "A2"

    ws_map = wb.create_sheet("Mapping_prestations")
    for col_idx, header in enumerate(["Code", "Libellé de prestation"], start=1):
        _make_header_cell(ws_map, 1, col_idx, header, hdr_font, fill_dark, border)
    ws_map.row_dimensions[1].height = 18
    for r_idx, (code_m, libelle_m) in enumerate(MAPPING_PRESTATIONS, start=2):
        cell_c = ws_map.cell(row=r_idx, column=1, value=code_m)
        cell_c.font = data_font
        cell_c.border = border
        cell_c.alignment = Alignment(horizontal="center", vertical="center")
        cell_l = ws_map.cell(row=r_idx, column=2, value=libelle_m)
        cell_l.font = data_font
        cell_l.border = border
        cell_l.alignment = Alignment(vertical="center")
    ws_map.column_dimensions["A"].width = 8
    ws_map.column_dimensions["B"].width = 70
    ws_map.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


def process_pdf(pdf_bytes: bytes, excel_bytes: bytes = None, site_key: str = None) -> tuple:
    """Returns (zip_bytes, stats_dict, unknown_codes_list)."""
    unknown_codes = []

    doc = fitz.open(stream=pdf_bytes, filetype="pdf")
    nb_pages = len(doc)

    debuts = [
        i for i in range(nb_pages)
        if re.search(r'Mission\s+d[ée]but[ée]e?\s+le', doc[i].get_text(), re.IGNORECASE)
    ]

    signes_pdfs = {}
    non_signes_pdfs = {}
    noms_utilises = {}

    for idx, debut in enumerate(debuts):
        fin = debuts[idx + 1] if idx + 1 < len(debuts) else nb_pages
        texte_premier = doc[debut].get_text()
        site, date_str, code = extraire_metadonnees_page(texte_premier, unknown_codes)

        nom_base = f"{site}_{code}_{date_str}.pdf"
        if nom_base in noms_utilises:
            noms_utilises[nom_base] += 1
            b, ext = os.path.splitext(nom_base)
            nom_final = f"{b}({noms_utilises[nom_base]}){ext}"
        else:
            noms_utilises[nom_base] = 0
            nom_final = nom_base

        bon_doc = fitz.open()
        for p in range(debut, fin):
            bon_doc.insert_pdf(doc, from_page=p, to_page=p)

        supprimer_note(bon_doc)
        signe = has_client_signature(bon_doc)

        buf = io.BytesIO()
        bon_doc.save(buf)
        bon_doc.close()

        if signe:
            signes_pdfs[nom_final] = buf.getvalue()
        else:
            non_signes_pdfs[nom_final] = buf.getvalue()

    doc.close()

    df_pointages = None
    if excel_bytes and site_key:
        site_nom = SITES_ALDI.get(site_key)
        if site_nom:
            df_pointages = charger_pointages_filtres(excel_bytes, site_nom)

    excel_out = generer_excel(signes_pdfs, non_signes_pdfs, df_pointages)

    zip_buf = io.BytesIO()
    with zipfile.ZipFile(zip_buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in signes_pdfs.items():
            zf.writestr(f"client_signe/{name}", data)
        for name, data in non_signes_pdfs.items():
            zf.writestr(f"client_non_signe/{name}", data)
        zf.writestr("recap_signes_non_signes.xlsx", excel_out)
        if unknown_codes:
            lines = "\n".join(f"- {c}" for c in unknown_codes)
            zf.writestr(
                "prestations_inconnues.txt",
                f"Prestations sans code reconnu (assignées UNK):\n\n{lines}\n"
            )

    zip_buf.seek(0)

    stats = {
        "total": len(debuts),
        "signes": len(signes_pdfs),
        "non_signes": len(non_signes_pdfs),
        "pointages": len(df_pointages) if df_pointages is not None else None,
        "unknowns": len(unknown_codes),
    }

    return zip_buf.read(), stats, unknown_codes
